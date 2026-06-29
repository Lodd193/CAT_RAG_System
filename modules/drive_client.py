import io
import json

import openpyxl

from docx import Document
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

import config

_SCOPES = ["https://www.googleapis.com/auth/drive"]
_service_cache = None


def _get_service():
    global _service_cache
    if _service_cache is None:
        creds_info = json.loads(config.GOOGLE_SERVICE_ACCOUNT_JSON)
        creds = service_account.Credentials.from_service_account_info(creds_info, scopes=_SCOPES)
        _service_cache = build("drive", "v3", credentials=creds)
    return _service_cache


def list_files(folder_id):
    """Return list of {id, name, mimeType} dicts for files in a Drive folder."""
    service = _get_service()
    query = f"'{folder_id}' in parents and trashed = false"
    result = service.files().list(q=query, fields="files(id, name, mimeType)").execute()
    return result.get("files", [])


def read_file(file_id, mime_type):
    """Return plain text content of a Drive file."""
    service = _get_service()

    if mime_type == "application/vnd.google-apps.document":
        response = service.files().export(fileId=file_id, mimeType="text/plain").execute()
        return response.decode("utf-8")

    request = service.files().get_media(fileId=file_id)
    buf = io.BytesIO()
    downloader = MediaIoBaseDownload(buf, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    buf.seek(0)

    if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(buf)
        return "\n".join(p.text for p in doc.paragraphs)

    if mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        wb = openpyxl.load_workbook(buf, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            parts.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                line = "\t".join("" if v is None else str(v) for v in row)
                if line.strip():
                    parts.append(line)
        return "\n".join(parts)

    return buf.read().decode("utf-8", errors="replace")


def get_folder_id(parent_folder_id, folder_name):
    """Return the Drive ID of a named subfolder within a parent folder."""
    service = _get_service()
    query = (
        f"'{parent_folder_id}' in parents"
        f" and name = '{folder_name}'"
        f" and mimeType = 'application/vnd.google-apps.folder'"
        f" and trashed = false"
    )
    result = service.files().list(q=query, fields="files(id, name)").execute()
    files = result.get("files", [])
    if not files:
        raise ValueError(f"Folder '{folder_name}' not found in {parent_folder_id}")
    return files[0]["id"]


def create_folder(parent_folder_id, folder_name):
    """Create a subfolder and return its ID. No-op if it already exists."""
    service = _get_service()
    query = (
        f"'{parent_folder_id}' in parents"
        f" and name = '{folder_name}'"
        f" and mimeType = 'application/vnd.google-apps.folder'"
        f" and trashed = false"
    )
    existing = service.files().list(q=query, fields="files(id)").execute().get("files", [])
    if existing:
        return existing[0]["id"]
    metadata = {
        "name": folder_name,
        "mimeType": "application/vnd.google-apps.folder",
        "parents": [parent_folder_id],
    }
    folder = service.files().create(body=metadata, fields="id").execute()
    return folder["id"]


def move_file(file_id, destination_folder_id):
    """Move a file to a different folder."""
    service = _get_service()
    file_meta = service.files().get(fileId=file_id, fields="parents").execute()
    previous_parents = ",".join(file_meta.get("parents", []))
    service.files().update(
        fileId=file_id,
        addParents=destination_folder_id,
        removeParents=previous_parents,
        fields="id, parents",
    ).execute()


def upload_file(folder_id, filename, content_bytes, mime_type):
    """Upload a file to a Drive folder. Returns the new file ID."""
    from googleapiclient.http import MediaInMemoryUpload
    service = _get_service()
    metadata = {"name": filename, "parents": [folder_id]}
    media = MediaInMemoryUpload(content_bytes, mimetype=mime_type)
    result = service.files().create(body=metadata, media_body=media, fields="id").execute()
    return result["id"]
